import csv
import json
from pathlib import Path
from datetime import datetime


HEADERS = [
    "ID", "Source", "Titre", "Entreprise", "Localisation",
    "Description", "URL", "Score", "Potentiel", "Secteur",
    "Taille", "Decideur", "Approche", "Points de contact",
    "Risques", "Opportunites", "Date scraping"
]

def _lead_to_row(lead):
    """Convertit un lead (dict) en ligne pour Excel/CSV."""
    enrichment = lead.get("ai_enrichment", {})
    # Support aussi format DB (flat)
    if not enrichment and lead.get("score_qualification") is not None:
        enrichment = {
            "score_qualification": lead.get("score_qualification", ""),
            "potentiel_estime": lead.get("potentiel_estime", ""),
            "secteur_activite": lead.get("secteur_activite", ""),
            "taille_entreprise": lead.get("taille_entreprise", ""),
            "decision_maker": bool(lead.get("decision_maker")),
            "approche_recommandee": lead.get("approche_recommandee", ""),
            "points_de_contact": _parse_json_field(lead.get("points_de_contact")),
            "risques": _parse_json_field(lead.get("risques")),
            "opportunites": _parse_json_field(lead.get("opportunites")),
        }

    def join_list(v):
        if isinstance(v, list):
            return "; ".join(v)
        return v or ""

    return [
        lead.get("id", ""),
        lead.get("source", ""),
        lead.get("title", lead.get("name", "")),
        lead.get("company", ""),
        lead.get("location", ""),
        lead.get("description", lead.get("headline", "")),
        lead.get("url", lead.get("profile_url", "")),
        enrichment.get("score_qualification", ""),
        enrichment.get("potentiel_estime", ""),
        enrichment.get("secteur_activite", ""),
        enrichment.get("taille_entreprise", ""),
        "Oui" if enrichment.get("decision_maker") else "Non" if enrichment.get("decision_maker") is not None and enrichment.get("decision_maker") != "" else "",
        enrichment.get("approche_recommandee", ""),
        join_list(enrichment.get("points_de_contact", "")),
        join_list(enrichment.get("risques", "")),
        join_list(enrichment.get("opportunites", "")),
        lead.get("scraped_at", lead.get("created_at", "")),
    ]

def _parse_json_field(value):
    if isinstance(value, str):
        try:
            return json.loads(value)
        except:
            return value
    return value or []

def export_to_csv(leads, filename=None):
    """Exporte les leads en CSV (gratuit, sans API)."""
    if filename is None:
        filename = f"data/leads_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    path = Path(filename)
    path.parent.mkdir(parents=True, exist_ok=True)

    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(HEADERS)
        for lead in leads:
            writer.writerow(_lead_to_row(lead))

    print(f"[CSV] {len(leads)} leads exportés -> {path}")
    return str(path)

def export_to_excel(leads, filename=None):
    """Exporte les leads en Excel .xlsx (gratuit, sans API)."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise ImportError("openpyxl non installé. Lance: pip install openpyxl")

    if filename is None:
        filename = f"data/leads_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = Path(filename)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"

    # Header style
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )

    # Écrire headers
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Coloration par score
    score_colors = {
        "hot": PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid"),
        "warm": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        "cold": PatternFill(start_color="DEEAF6", end_color="DEEAF6", fill_type="solid"),
    }

    for row_idx, lead in enumerate(leads, 2):
        row = _lead_to_row(lead)
        score = str(row[7]).lower().strip()
        fill = score_colors.get(score)

        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if fill:
                cell.fill = fill
            # Score en gras + couleur texte
            if col_idx == 8 and score:
                if score == "hot":
                    cell.font = Font(bold=True, color="C0392B")
                elif score == "warm":
                    cell.font = Font(bold=True, color="D48806")
                elif score == "cold":
                    cell.font = Font(bold=True, color="2471A3")

    # Largeurs colonnes
    widths = [6, 12, 22, 18, 16, 35, 30, 10, 10, 16, 12, 10, 30, 25, 25, 25, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(path)
    print(f"[Excel] {len(leads)} leads exportés -> {path}")
    return str(path)

def export_to_json(leads, filename=None):
    """Exporte en JSON (pour debug/API)."""
    if filename is None:
        filename = f"data/leads_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    path = Path(filename)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(leads, f, indent=2, ensure_ascii=False)
    print(f"[JSON] {len(leads)} leads exportés -> {path}")
    return str(path)
